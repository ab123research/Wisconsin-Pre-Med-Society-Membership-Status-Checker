"""Read the membership workbook and check one member. Also usable as a CLI.

The spreadsheet's Active Member field is authoritative. Calculated requirements
describe progress only; they never override that field or apply exceptions.
"""

from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

DEFAULT_SHEET_ID = "1fp9zse-tPuMC2zBfiUiV9WW5RnR8uA38lqDcDSgSoo8"
DEFAULT_WORKSHEET = "Fall 26 membership"
REQUIRED_MEETINGS = 3
MAX_DOWNLOAD_BYTES = 10 * 1024 * 1024
EVENT_HEADERS = {
    "volunteering": "Volunteered?",
    "social": "Social Event Attended?",
    "mentorship": "Mentorship Program?",
    "office_hours": "Attended Office Hours?",
    "workshop": "Workshop Attended?",
    "fundraising": "Fundraised?",
}
EVENT_LABELS = {
    "volunteering": "Volunteering",
    "social": "Social event",
    "mentorship": "Mentorship program",
    "office_hours": "Office hours",
    "workshop": "Workshop",
    "fundraising": "Fundraising",
}


class MembershipError(Exception):
    """A readable source-data or lookup error."""


class MemberNotFound(MembershipError):
    pass


class DuplicateMember(MembershipError):
    pass


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.translate(str.maketrans({"’": "'", "‘": "'", "–": "-", "—": "-"}))
    return " ".join(value.split()).casefold()


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize(str(value or "")))


def parse_flag(value: object, *, blank_is_false: bool = True) -> bool | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return False if blank_is_false else None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return {0: False, 1: True}.get(value)
    if isinstance(value, str):
        text = normalize(value)
        if text in {"x", "yes", "y", "true", "1", "checked", "attended", "✓", "✔"}:
            return True
        if text in {"no", "n", "false", "0", "unchecked", "not attended"}:
            return False
    return None


def parse_dues(value: object) -> bool | None:
    if isinstance(value, str):
        text = normalize(value)
        if text in {"year", "semester", "paid", "fall", "fall 2026"}:
            return True
        if text in {"unpaid", "not paid"}:
            return False
    return parse_flag(value, blank_is_false=False)


def any_attended(flags: list[bool | None]) -> bool | None:
    if any(flag is True for flag in flags):
        return True
    return None if any(flag is None for flag in flags) else False


@dataclass(frozen=True)
class Activity:
    label: str
    attended: bool | None


@dataclass(frozen=True)
class Member:
    name: str
    row_number: int
    official_active: bool | None
    dues_paid: bool | None
    dues_record: str
    meetings: tuple[Activity, ...]
    events: dict[str, Activity]
    issues: tuple[str, ...]

    @property
    def meetings_attended(self) -> int:
        return sum(meeting.attended is True for meeting in self.meetings)

    @property
    def meeting_requirement(self) -> bool | None:
        if self.meetings_attended >= REQUIRED_MEETINGS:
            return True
        if any(meeting.attended is None for meeting in self.meetings):
            return None
        return False

    @property
    def purple_met(self) -> bool | None:
        return any_attended([self.events[key].attended for key in ("volunteering", "social")])

    @property
    def pink_met(self) -> bool | None:
        return any_attended([self.events[key].attended for key in ("mentorship", "office_hours", "workshop")])

    @property
    def requirement_states(self) -> tuple[bool | None, ...]:
        return (self.dues_paid, self.meeting_requirement, self.purple_met,
                self.pink_met, self.events["fundraising"].attended)

    @property
    def requirements_met(self) -> int:
        return sum(state is True for state in self.requirement_states)

    @property
    def remaining(self) -> list[str]:
        items = []
        if self.dues_paid is not True:
            items.append("Pay semester dues." if self.dues_paid is False else "Ask the secretary to confirm your dues payment.")
        if self.meeting_requirement is False:
            needed = REQUIRED_MEETINGS - self.meetings_attended
            items.append(f"Attend {needed} more general meeting{'s' if needed != 1 else ''}.")
        elif self.meeting_requirement is None:
            items.append("Ask the secretary to confirm the unrecognized general-meeting attendance entries.")
        for state, action, unknown in (
            (self.purple_met, "Attend a volunteering or social event (purple).", "volunteering/social attendance"),
            (self.pink_met, "Participate in mentorship, office hours, or a workshop (pink).", "mentorship/office hours/workshop attendance"),
            (self.events["fundraising"].attended, "Participate in fundraising (orange).", "fundraising attendance"),
        ):
            if state is False:
                items.append(action)
            elif state is None:
                items.append(f"Ask the secretary to confirm your {unknown}.")
        return items

    def report(self) -> dict:
        return {**asdict(self), "meetings_attended": self.meetings_attended,
                "meetings_total": len(self.meetings),
                "required_meetings": REQUIRED_MEETINGS,
                "purple_met": self.purple_met, "pink_met": self.pink_met,
                "requirements_met": self.requirements_met, "remaining": self.remaining}


@dataclass(frozen=True)
class MembershipBook:
    members: tuple[Member, ...]
    worksheet: str
    fetched_at: datetime

    def find(self, name: str) -> Member:
        query = normalize(name)
        if not query:
            raise MemberNotFound("Enter your full name as it appears on the membership sheet.")
        matches = [member for member in self.members if normalize(member.name) == query]
        if not matches:
            raise MemberNotFound("No matching member was found. Check your full name and spelling, or contact the secretary.")
        if len(matches) > 1:
            raise DuplicateMember("More than one record has this name. Contact the secretary to identify the correct record.")
        return matches[0]


def download_workbook(sheet_id: str = DEFAULT_SHEET_ID) -> bytes:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", sheet_id):
        raise MembershipError("The configured Google Sheet ID is invalid.")
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    request = Request(url, headers={"User-Agent": "MembershipChecker/1.0"})
    try:
        with urlopen(request, timeout=20) as response:
            data = response.read(MAX_DOWNLOAD_BYTES + 1)
    except HTTPError as exc:
        if exc.code in {401, 403, 404}:
            raise MembershipError("The membership sheet is unavailable. The secretary should check its public viewing and download permissions.") from exc
        raise MembershipError("Google Sheets could not provide the membership data. Please try again shortly.") from exc
    except (URLError, TimeoutError, OSError) as exc:
        raise MembershipError("The membership sheet could not be reached. Please try again shortly.") from exc
    if len(data) > MAX_DOWNLOAD_BYTES:
        raise MembershipError("The membership workbook exceeds the 10 MB download limit.")
    if not data.startswith(b"PK"):
        raise MembershipError("Google Sheets did not return an Excel file. The secretary should check its public viewing and download permissions.")
    return data


def read_workbook(data: bytes, worksheet: str = DEFAULT_WORKSHEET) -> MembershipBook:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError) as exc:
        raise MembershipError("The membership workbook could not be read as a valid Excel file.") from exc
    try:
        if worksheet not in workbook.sheetnames:
            raise MembershipError(f'The configured worksheet "{worksheet}" is missing. Ask the secretary to check the app settings.')
        sheet = workbook[worksheet]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        columns = {}
        for index, value in enumerate(headers):
            if value is None or not str(value).strip():
                continue
            key = header_key(value)
            if key in columns:
                raise MembershipError("The membership sheet has duplicate column headers. Ask the secretary to correct them.")
            columns[key] = index

        def column(label: str) -> int:
            key = header_key(label)
            if key not in columns:
                raise MembershipError(f'The membership sheet is missing the "{label}" column.')
            return columns[key]

        name_col = column("Name")
        active_col = column("Active Member?")
        dues_col = column("Dues Paid?")
        total_col = column("Total Meetings Attended (3)")
        meeting_cols = [index for index in range(dues_col + 1, total_col)
                        if headers[index] is not None and str(headers[index]).strip()]
        if not meeting_cols:
            raise MembershipError("No general-meeting columns were found between dues and the meeting total.")
        event_cols = {key: column(label) for key, label in EVENT_HEADERS.items()}
        reserved = {name_col, active_col, dues_col, total_col, *event_cols.values()}
        if reserved.intersection(meeting_cols):
            raise MembershipError("The membership column layout changed. Ask the secretary to review the app configuration.")
        members = []
        for row_number, row in enumerate(rows, start=2):
            name = row[name_col]
            if not isinstance(name, str) or not name.strip():
                continue
            if normalize(name).rstrip(":") in {"total", "totals", "grand total", "grand totals"}:
                continue
            issues = []

            def flag(index: int, *, blank_is_false: bool = True) -> bool | None:
                value = parse_flag(row[index], blank_is_false=blank_is_false)
                if value is None:
                    issues.append(f'Unrecognized or missing value in "{str(headers[index]).replace(chr(10), " ")}".')
                return value

            official_active = flag(active_col, blank_is_false=False)
            dues_paid = parse_dues(row[dues_col])
            if dues_paid is None:
                issues.append('Unrecognized or missing value in "Dues Paid?".')
            meetings = tuple(Activity(str(headers[index]).strip(), flag(index)) for index in meeting_cols)
            events = {key: Activity(EVENT_LABELS[key], flag(index)) for key, index in event_cols.items()}
            members.append(Member(name.strip(), row_number, official_active, dues_paid,
                                  str(row[dues_col] or "Not recorded"), meetings, events, tuple(issues)))
        if not members:
            raise MembershipError("The membership sheet does not contain any named member records.")
        return MembershipBook(tuple(members), worksheet, datetime.now(timezone.utc))
    finally:
        workbook.close()


def load_membership(*, file: str | None = None, sheet_id: str = DEFAULT_SHEET_ID,
                    worksheet: str = DEFAULT_WORKSHEET) -> MembershipBook:
    if file:
        try:
            data = Path(file).read_bytes()
        except OSError as exc:
            raise MembershipError("The local membership workbook could not be opened.") from exc
    else:
        data = download_workbook(sheet_id)
    return read_workbook(data, worksheet)


def state_text(state: bool | None, yes: str, no: str) -> str:
    return "Not recorded / needs review" if state is None else yes if state else no


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--name", help="Member's full name; prompts if omitted")
    parser.add_argument("--file", help="Optional local Excel workbook instead of the live Google Sheet")
    parser.add_argument("--sheet-id", default=DEFAULT_SHEET_ID)
    parser.add_argument("--worksheet", default=DEFAULT_WORKSHEET)
    parser.add_argument("--json", action="store_true", help="Output a JSON report")
    args = parser.parse_args()
    try:
        name = args.name if args.name is not None else input("Full name: ")
        member = load_membership(file=args.file, sheet_id=args.sheet_id, worksheet=args.worksheet).find(name)
    except (MembershipError, EOFError) as exc:
        if args.json:
            print(json.dumps({"error": str(exc)}))
        else:
            print(f"Unable to check membership: {exc}")
        return 1
    if args.json:
        print(json.dumps(member.report(), indent=2, ensure_ascii=False))
        return 0
    print(f"\n{member.name}")
    print("Official membership status: " + state_text(member.official_active, "Active", "Inactive"))
    print("Semester dues: " + state_text(member.dues_paid, "Paid", "Not paid"))
    print(f"General meetings: {member.meetings_attended} / {len(member.meetings)} (at least {REQUIRED_MEETINGS} required)")
    for meeting in member.meetings:
        print(f"  {meeting.label}: {state_text(meeting.attended, 'Attended', 'Not Attended')}")
    for event in member.events.values():
        print(f"{event.label}: {state_text(event.attended, 'Attended', 'Not Attended')}")
    print(f"\nRequirements met: {member.requirements_met} / 5")
    for item in member.remaining:
        print(f"- {item}")
    if not member.remaining:
        print("All listed requirements are met.")
        if member.official_active is not True:
            print("Ask the secretary to review your official membership status.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
