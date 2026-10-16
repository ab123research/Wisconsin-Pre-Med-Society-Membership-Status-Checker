"""Read the membership workbook and check one member. Also usable as a CLI.

The spreadsheet's Active Member field is authoritative. Calculated requirements
describe progress only; they never override that field or apply exceptions.
"""

from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata
from collections.abc import Mapping
import tomllib
from urllib.parse import quote

import requests
from google.auth.exceptions import GoogleAuthError
from google.auth.transport.requests import AuthorizedSession
from google.oauth2.service_account import Credentials as GoogleCredentials
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

DEFAULT_WORKSHEET = "Fall 26 membership"
REQUIRED_MEETINGS = 3
SHEETS_READONLY_SCOPES = ("https://www.googleapis.com/auth/spreadsheets.readonly",)
EVENT_HEADERS = {
    "volunteering": "Volunteered?",
    "social": "Social Event Attended?",
    "mentorship": "Mentorship Program?",
    "office_hours": "Attended Office Hours?",
    "workshop": "Workshop Attended?",
    "fundraising": "Fundraised?",
}
EVENT_LABELS = {
    "volunteering": "Volunteering",
    "social": "Social event",
    "mentorship": "Mentorship program",
    "office_hours": "Office hours",
    "workshop": "Workshop",
    "fundraising": "Fundraising",
}


class MembershipError(Exception):
    """A readable source-data or lookup error."""


class MemberNotFound(MembershipError):
    pass


class DuplicateMember(MembershipError):
    pass


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.translate(str.maketrans({"’": "'", "‘": "'", "–": "-", "—": "-"}))
    return " ".join(value.split()).casefold()


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize(str(value or "")))


def parse_flag(value: object, *, blank_is_false: bool = True) -> bool | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return False if blank_is_false else None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return {0: False, 1: True}.get(value)
    if isinstance(value, str):
        text = normalize(value)
        if text in {"x", "yes", "y", "true", "1", "checked", "attended", "✓", "✔"}:
            return True
        if text in {"no", "n", "false", "0", "unchecked", "not attended"}:
            return False
    return None


def parse_dues(value: object) -> bool | None:
    if isinstance(value, str):
        text = normalize(value)
        if text in {"year", "semester", "paid", "fall", "fall 2026"}:
            return True
        if text in {"unpaid", "not paid"}:
            return False
    return parse_flag(value, blank_is_false=False)


def any_attended(flags: list[bool | None]) -> bool | None:
    if any(flag is True for flag in flags):
        return True
    return None if any(flag is None for flag in flags) else False


@dataclass(frozen=True)
class Activity:
    label: str
    attended: bool | None


@dataclass(frozen=True)
class Member:
    name: str
    row_number: int
    official_active: bool | None
    dues_paid: bool | None
    dues_record: str
    meetings: tuple[Activity, ...]
    events: dict[str, Activity]
    issues: tuple[str, ...]

    @property
    def meetings_attended(self) -> int:
        return sum(meeting.attended is True for meeting in self.meetings)

    @property
    def meeting_requirement(self) -> bool | None:
        if self.meetings_attended >= REQUIRED_MEETINGS:
            return True
        if any(meeting.attended is None for meeting in self.meetings):
            return None
        return False

    @property
    def purple_met(self) -> bool | None:
        return any_attended([self.events[key].attended for key in ("volunteering", "social")])

    @property
    def pink_met(self) -> bool | None:
        return any_attended([self.events[key].attended for key in ("mentorship", "office_hours", "workshop")])

    @property
    def requirement_states(self) -> tuple[bool | None, ...]:
        return (self.dues_paid, self.meeting_requirement, self.purple_met,
                self.pink_met, self.events["fundraising"].attended)

    @property
    def requirements_met(self) -> int:
        return sum(state is True for state in self.requirement_states)

    @property
    def remaining(self) -> list[str]:
        items = []
        if self.dues_paid is not True:
            items.append("Pay semester dues." if self.dues_paid is False else "Ask the secretary to confirm your dues payment.")
        if self.meeting_requirement is False:
            needed = REQUIRED_MEETINGS - self.meetings_attended
            items.append(f"Attend {needed} more general meeting{'s' if needed != 1 else ''}.")
        elif self.meeting_requirement is None:
            items.append("Ask the secretary to confirm the unrecognized general-meeting attendance entries.")
        for state, action, unknown in (
            (self.purple_met, "Attend a volunteering or social event (purple).", "volunteering/social attendance"),
            (self.pink_met, "Participate in mentorship, office hours, or a workshop (pink).", "mentorship/office hours/workshop attendance"),
            (self.events["fundraising"].attended, "Participate in fundraising (orange).", "fundraising attendance"),
        ):
            if state is False:
                items.append(action)
            elif state is None:
                items.append(f"Ask the secretary to confirm your {unknown}.")
        return items

    def report(self) -> dict:
        return {**asdict(self), "meetings_attended": self.meetings_attended,
                "meetings_total": len(self.meetings),
                "required_meetings": REQUIRED_MEETINGS,
                "purple_met": self.purple_met, "pink_met": self.pink_met,
                "requirements_met": self.requirements_met, "remaining": self.remaining}


@dataclass(frozen=True)
class MembershipBook:
    members: tuple[Member, ...]
    worksheet: str
    fetched_at: datetime

    def find(self, name: str) -> Member:
        query = normalize(name)
        if not query:
            raise MemberNotFound("Enter your full name as it appears on the membership sheet.")
        matches = [member for member in self.members if normalize(member.name) == query]
        if not matches:
            raise MemberNotFound("No matching member was found. Check your full name and spelling, or contact the secretary.")
        if len(matches) > 1:
            raise DuplicateMember("More than one record has this name. Contact the secretary to identify the correct record.")
        return matches[0]


@dataclass(frozen=True)
class ConnectionSettings:
    spreadsheet_id: str
    worksheet: str
    service_account_info: dict = field(repr=False)


def parse_connection_settings(section: Mapping) -> ConnectionSettings:
    """Parse the documented [google_sheets] section without echoing secrets."""
    if not isinstance(section, Mapping):
        raise MembershipError("The google_sheets secrets section is missing or invalid.")
    sheet_id = section.get("spreadsheet_id", "")
    if not isinstance(sheet_id, str) or not re.fullmatch(r"[A-Za-z0-9_-]+", sheet_id):
        raise MembershipError("Set spreadsheet_id to the sheet ID in the google_sheets secrets section.")
    worksheet = section.get("worksheet", DEFAULT_WORKSHEET)
    if not isinstance(worksheet, str) or not worksheet.strip():
        raise MembershipError("Set worksheet to the membership tab name in the google_sheets secrets section.")
    raw = section.get("service_account_json")
    if not isinstance(raw, str):
        raise MembershipError("Set service_account_json to the entire downloaded JSON key inside triple single quotes.")
    try:
        info = json.loads(raw)
    except (ValueError, TypeError):
        raise MembershipError("service_account_json is not valid JSON. Copy the entire downloaded key, including both braces, between the triple single quotes.") from None
    if (not isinstance(info, dict) or info.get("type") != "service_account"
            or not all(isinstance(info.get(key), str) and info[key].strip()
                       for key in ("client_email", "private_key", "token_uri"))):
        raise MembershipError("service_account_json must contain the complete downloaded service-account key.")
    if info["token_uri"] not in {"https://oauth2.googleapis.com/token", "https://accounts.google.com/o/oauth2/token"}:
        raise MembershipError("The key's authentication endpoint is not a supported Google endpoint.")
    return ConnectionSettings(sheet_id, worksheet, info)


def read_secrets_file(path: str | Path) -> ConnectionSettings:
    try:
        with Path(path).open("rb") as source:
            config = tomllib.load(source)
    except (OSError, ValueError):
        # TOML exceptions can include credential text; never expose their text.
        raise MembershipError("The local secrets file could not be read. Check .streamlit/secrets.toml and its TOML formatting.") from None
    return parse_connection_settings(config.get("google_sheets"))


def read_private_sheet(sheet_id: str, worksheet: str, service_account_info: dict | None) -> MembershipBook:
    """Read a private Google Sheet with a scoped service account. No public fallback."""
    if not isinstance(sheet_id, str) or not re.fullmatch(r"[A-Za-z0-9_-]+", sheet_id):
        raise MembershipError("The configured Google Sheet ID is invalid or missing.")
    if not service_account_info:
        raise MembershipError("Private-sheet credentials are missing. Configure the google_sheets secrets section.")
    try:
        credentials = GoogleCredentials.from_service_account_info(
            service_account_info, scopes=SHEETS_READONLY_SCOPES)
    except (ValueError, TypeError, KeyError, GoogleAuthError):
        raise MembershipError("The service-account key could not be loaded. Check that service_account_json contains the full, valid downloaded key.") from None
    # Quoting the whole tab name prevents special characters from changing the range.
    sheet_range = "'" + worksheet.replace("'", "''") + "'"
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_id}/values/{quote(sheet_range, safe='')}"
    try:
        with AuthorizedSession(credentials, refresh_timeout=20) as session:
            response = session.get(url, params={"majorDimension": "ROWS",
                "valueRenderOption": "UNFORMATTED_VALUE"}, timeout=20)
            if response.status_code == 400:
                raise MembershipError("Google could not read the configured worksheet. Check the worksheet name in your secrets.")
            if response.status_code in {401, 403, 404}:
                raise MembershipError("The private membership sheet could not be accessed. Check the sheet ID, Viewer sharing with the service account, and that Google Sheets API is enabled.")
            if response.status_code != 200:
                raise MembershipError("Google Sheets is temporarily unavailable. Please try again shortly.")
            payload = response.json()
    except GoogleAuthError:
        raise MembershipError("Google authentication failed. Check that the service-account key is active and copied correctly in your secrets.") from None
    except requests.RequestException:
        raise MembershipError("The private membership sheet could not be reached. Please try again shortly.") from None
    except (ValueError, TypeError):
        raise MembershipError("Google Sheets returned an unreadable response. Please try again shortly.") from None
    if not isinstance(payload, dict) or not isinstance(payload.get("values", []), list):
        raise MembershipError("Google Sheets returned an unexpected data format.")
    values = payload.get("values", [])
    if any(not isinstance(row, list) for row in values):
        raise MembershipError("Google Sheets returned an unexpected row format.")
    return read_rows(values, worksheet)


def read_workbook(data: bytes, worksheet: str = DEFAULT_WORKSHEET) -> MembershipBook:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError) as exc:
        raise MembershipError("The membership workbook could not be read as a valid Excel file.") from exc
    try:
        if worksheet not in workbook.sheetnames:
            raise MembershipError(f'The configured worksheet "{worksheet}" is missing. Ask the secretary to check the app settings.')
        sheet = workbook[worksheet]
        return read_rows(sheet.iter_rows(values_only=True), worksheet)
    finally:
        workbook.close()


def read_rows(rows, worksheet: str = DEFAULT_WORKSHEET) -> MembershipBook:
    """Interpret the same header/row layout from Excel or authenticated Sheets API."""
    rows = iter(rows)
    headers = next(rows, ())
    columns = {}
    for index, value in enumerate(headers):
        if value is None or not str(value).strip():
            continue
        key = header_key(value)
        if key in columns:
            raise MembershipError("The membership sheet has duplicate column headers. Ask the secretary to correct them.")
        columns[key] = index

    def column(label: str) -> int:
        key = header_key(label)
        if key not in columns:
            raise MembershipError(f'The membership sheet is missing the "{label}" column.')
        return columns[key]

    name_col = column("Name")
    active_col = column("Active Member?")
    dues_col = column("Dues Paid?")
    total_col = column("Total Meetings Attended (3)")
    meeting_cols = [index for index in range(dues_col + 1, total_col)
                    if headers[index] is not None and str(headers[index]).strip()]
    if not meeting_cols:
        raise MembershipError("No general-meeting columns were found between dues and the meeting total.")
    event_cols = {key: column(label) for key, label in EVENT_HEADERS.items()}
    reserved = {name_col, active_col, dues_col, total_col, *event_cols.values()}
    if reserved.intersection(meeting_cols):
        raise MembershipError("The membership column layout changed. Ask the secretary to review the app configuration.")
    members = []
    for row_number, raw_row in enumerate(rows, start=2):
        row = tuple(raw_row) + (None,) * max(0, len(headers) - len(raw_row))
        name = row[name_col]
        if not isinstance(name, str) or not name.strip():
            continue
        if normalize(name).rstrip(":") in {"total", "totals", "grand total", "grand totals"}:
            continue
        issues = []

        def flag(index: int, *, blank_is_false: bool = True) -> bool | None:
            value = parse_flag(row[index], blank_is_false=blank_is_false)
            if value is None:
                issues.append(f'Unrecognized or missing value in "{str(headers[index]).replace(chr(10), " ")}".')
            return value

        official_active = flag(active_col, blank_is_false=False)
        dues_paid = parse_dues(row[dues_col])
        if dues_paid is None:
            issues.append('Unrecognized or missing value in "Dues Paid?".')
        meetings = tuple(Activity(str(headers[index]).strip(), flag(index)) for index in meeting_cols)
        events = {key: Activity(EVENT_LABELS[key], flag(index)) for key, index in event_cols.items()}
        members.append(Member(name.strip(), row_number, official_active, dues_paid,
                              str(row[dues_col] or "Not recorded"), meetings, events, tuple(issues)))
    if not members:
        raise MembershipError("The membership sheet does not contain any named member records.")
    return MembershipBook(tuple(members), worksheet, datetime.now(timezone.utc))


def load_membership(*, file: str | None = None, sheet_id: str = "",
                    worksheet: str = DEFAULT_WORKSHEET,
                    service_account_info: dict | None = None) -> MembershipBook:
    if file:
        try:
            data = Path(file).read_bytes()
        except OSError:
            raise MembershipError("The local membership workbook could not be opened.") from None
        return read_workbook(data, worksheet)
    return read_private_sheet(sheet_id, worksheet, service_account_info)


def state_text(state: bool | None, yes: str, no: str) -> str:
    return "Not recorded / needs review" if state is None else yes if state else no


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--name", help="Member's full name; prompts if omitted")
    parser.add_argument("--file", help="Optional local Excel workbook instead of the live Google Sheet")
    parser.add_argument("--sheet-id", help="Optional override for the ID in your secrets file")
    parser.add_argument("--secrets-file", default=str(Path(__file__).resolve().parent / ".streamlit" / "secrets.toml"),
                        help="Private connection settings; defaults to .streamlit/secrets.toml beside this script")
    parser.add_argument("--worksheet", help="Optional worksheet name override")
    parser.add_argument("--json", action="store_true", help="Output a JSON report")
    args = parser.parse_args()
    try:
        name = args.name if args.name is not None else input("Full name: ")
        if args.file:
            book = load_membership(file=args.file, worksheet=args.worksheet or DEFAULT_WORKSHEET)
        else:
            settings = read_secrets_file(args.secrets_file)
            book = load_membership(sheet_id=args.sheet_id or settings.spreadsheet_id,
                                   worksheet=args.worksheet or settings.worksheet,
                                   service_account_info=settings.service_account_info)
        member = book.find(name)
    except (MembershipError, EOFError) as exc:
        if args.json:
            print(json.dumps({"error": str(exc)}))
        else:
            print(f"Unable to check membership: {exc}")
        return 1
    if args.json:
        print(json.dumps(member.report(), indent=2, ensure_ascii=False))
        return 0
    print(f"\n{member.name}")
    print("Official membership status: " + state_text(member.official_active, "Active", "Inactive"))
    print("Semester dues: " + state_text(member.dues_paid, "Paid", "Not paid"))
    print(f"General meetings: {member.meetings_attended} / {len(member.meetings)} (at least {REQUIRED_MEETINGS} required)")
    for meeting in member.meetings:
        print(f"  {meeting.label}: {state_text(meeting.attended, 'Attended', 'Not Attended')}")
    for event in member.events.values():
        print(f"{event.label}: {state_text(event.attended, 'Attended', 'Not Attended')}")
    print(f"\nRequirements met: {member.requirements_met} / 5")
    for item in member.remaining:
        print(f"- {item}")
    if not member.remaining:
        print("All listed requirements are met.")
        if member.official_active is not True:
            print("Ask the secretary to review your official membership status.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
